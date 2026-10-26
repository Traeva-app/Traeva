"""Creates isolated synthetic PowerPoint inputs; never reads user documents."""
from datetime import datetime, timezone
from pathlib import Path
import io
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml import parse_xml
from PIL import Image

OUT = Path(__file__).resolve().parent.parent / 'fixtures' / 'pptx'
OUT.mkdir(parents=True, exist_ok=True)
for rich in (False, True):
    deck = Presentation()
    deck.core_properties.author = 'Synthetic Reader Test'
    deck.core_properties.last_modified_by = 'Synthetic Reader Test'
    deck.core_properties.title = 'TRAEVA PPTX Reader Synthetic Fixture'
    deck.core_properties.created = datetime(2026, 10, 11, tzinfo=timezone.utc)
    deck.core_properties.modified = datetime(2026, 10, 11, tzinfo=timezone.utc)
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = 'First title 中国'
    first.placeholders[1].text = 'First body\nSecond paragraph 🦀'
    first.notes_slide.notes_text_frame.text = 'First speaker note\nNote line two'
    second = deck.slides.add_slide(deck.slide_layouts[1])
    second.shapes.title.text = 'Second title'
    second.placeholders[1].text = 'Second body'
    second.notes_slide.notes_text_frame.text = 'Second speaker note'
    second._element.set('show', '0')
    # Order differs from part filenames; notes must follow their own slide.
    ids = deck.slides._sldIdLst
    ids.insert(0, ids[-1])
    if rich:
        picture = io.BytesIO()
        Image.new('RGB', (16, 16), 'blue').save(picture, format='PNG')
        first.shapes.add_picture(picture, Inches(7), Inches(1), Inches(1), Inches(1))
        chart = CategoryChartData()
        chart.categories = ['Synthetic A', 'Synthetic B']
        chart.add_series('Synthetic series', (2, 5))
        first.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(3), Inches(2), chart)
        table = first.shapes.add_table(1, 1, Inches(5), Inches(4), Inches(2), Inches(1)).table
        table.cell(0, 0).text = 'OMITTED TABLE SECRET'
        transition = OxmlElement('p:transition')
        transition.append(OxmlElement('p:fade'))
        first._element.append(transition)
        # Own synthetic title Fade saved by PowerPoint 16.78; shape id 2.
        timing = parse_xml((OUT / 'powerpoint-animation.xml').read_bytes())
        first._element.append(timing)
    deck.save(OUT / ('rich-source.pptx' if rich else 'basic-source.pptx'))
expected = '[Slide 1]\n[Title]\nSecond title\n[Text]\nSecond body\n[Notes]\nSecond speaker note\n\n[Slide 2]\n[Title]\nFirst title 中国\n[Text]\nFirst body\nSecond paragraph 🦀\n[Notes]\nFirst speaker note\nNote line two'
(OUT / 'expected.txt').write_text(expected)
