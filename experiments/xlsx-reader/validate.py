"""Run offline checks and retain exact synthetic fixture/probe evidence."""
import datetime
import hashlib
import importlib.util
import json
from pathlib import Path
import platform
import subprocess
import sys
import warnings
import zipfile
import xml.etree.ElementTree as ET

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
LOCAL = HERE / "evidence" / "local"
LOCAL.mkdir(parents=True, exist_ok=True)


def command(args, name):
    p = subprocess.run(args, cwd=HERE, capture_output=True, text=True, timeout=120)
    (LOCAL / f"{name}.txt").write_text(p.stdout + p.stderr)
    if p.returncode:
        raise RuntimeError(f"{name} failed, see {LOCAL / (name + '.txt')}")
    return p.stdout


def probe(*names, expected_success=True):
    files = [HERE / "fixtures" / "generated" / n for n in names]
    initial = [hashlib.sha256(p.read_bytes()).hexdigest() for p in files]
    p = subprocess.run([str(HERE / "target" / "debug" / "xlsx-probe"), *map(str, files)],
                       cwd=HERE, capture_output=True, text=True, timeout=10)
    if (p.returncode == 0) != expected_success:
        raise RuntimeError(f"unexpected probe exit {p.returncode}: {names}: {p.stdout} {p.stderr}")
    result = json.loads(p.stdout)
    if not expected_success and "workbook" in result:
        raise RuntimeError("failed probe leaked a successful workbook")
    if initial != [hashlib.sha256(p.read_bytes()).hexdigest() for p in files]:
        raise RuntimeError("probe modified a fixture")
    (LOCAL / ("__".join(names) + ".json")).write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    return result


def main():
    spec = importlib.util.spec_from_file_location("fixture_builder", HERE / "fixtures" / "generate.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    fixtures = module.generate()
    checks = []
    for name, args in [
        ("fmt", ["cargo", "fmt", "--check"]),
        ("clippy", ["cargo", "clippy", "--offline", "--all-targets", "--", "-D", "warnings"]),
        ("tests", ["cargo", "test", "--offline"]),
        ("build", ["cargo", "build", "--offline", "--bin", "xlsx-probe"]),
    ]:
        text = command(args, name)
        checks.append({"name": name, "command": args, "status": "PASS"})
        if name == "tests":
            checks[-1]["test_summary"] = [s for s in text.splitlines() if "test result:" in s]

    results = []
    for name, case in fixtures.items():
        success = case["expected"] in ("partial", "complete_within_profile")
        data = probe(name, expected_success=success)
        actual = data["workbook"]["status"] if success else data["error"]["kind"]
        assert actual == case["expected"], (name, actual, case["expected"])
        if success:
            assert data["workbook"]["input_sha256"] == case["sha256"]
        results.append({"file": name, **case, "actual": actual, "status": "PASS"})
        if "xml_syntax_invalid_part" in case:
            with zipfile.ZipFile(HERE / "fixtures" / "generated" / name) as z:
                assert z.testzip() is None
                try:
                    ET.fromstring(z.read(case["xml_syntax_invalid_part"]))
                except ET.ParseError:
                    results[-1]["xml_syntax_cross_check"] = "REJECTED_BY_ELEMENTTREE"
                else:
                    raise RuntimeError(f"malformed XML fixture unexpectedly parsed: {name}")
        if success:
            with zipfile.ZipFile(HERE / "fixtures" / "generated" / name) as z:
                assert z.testzip() is None
                for part in z.namelist():
                    if part.endswith((".xml", ".rels")):
                        ET.fromstring(z.read(part))

    before = probe("before.xlsx")["workbook"]
    cells = {c["address"]: c for c in before["sheets"][0]["cells"]}
    assert cells["A2"]["literal"] == {"type": "text", "value": "Hosting"}
    assert cells["B1"]["literal"] == {"type": "text", "value": "USD"}
    assert cells["D2"]["literal"] == {"type": "number", "value": "9007199254740993"}
    assert cells["C2"]["formula"]["text"] == "B2*2"
    assert cells["C2"]["cached_result"] == {"value": {"type": "number", "value": "40"}, "freshness": "unverified"}
    assert cells["G2"]["cached_result"]["value"] is None
    assert before["sheets"][1]["visibility"] == "hidden"
    comparisons = []
    for name in ("after.xlsx", "cache-only.xlsx", "formula-only.xlsx"):
        diff = probe("before.xlsx", name)["coordinate_changes"]
        by_address = {c["address"]: c for c in diff}
        if name == "after.xlsx":
            assert set(by_address) == {"B2", "C2"}
            assert by_address["B2"]["literal_changed"]
            assert by_address["C2"]["cache_changed"] and not by_address["C2"]["formula_changed"]
        elif name == "cache-only.xlsx":
            assert set(by_address) == {"C2"}
            assert by_address["C2"]["cache_changed"]
            assert not by_address["C2"]["formula_changed"] and not by_address["C2"]["literal_changed"]
        else:
            assert set(by_address) == {"C2"}
            assert by_address["C2"]["formula_changed"] and not by_address["C2"]["cache_changed"]
        comparisons.append({"before": "before.xlsx", "after": name, "status": "PASS", "changes": diff})

    cross_reader = {"status": "NOT_RUN", "reason": "openpyxl unavailable in selected Python"}
    try:
        import openpyxl
    except ImportError:
        pass
    else:
        path = HERE / "fixtures" / "generated" / "before.xlsx"
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            formula = openpyxl.load_workbook(path, data_only=False, read_only=True, keep_links=False)
            saved = openpyxl.load_workbook(path, data_only=True, read_only=True, keep_links=False)
        assert formula.sheetnames == ["Budget", "Notes"]
        assert formula["Budget"]["C2"].value == "=B2*2"
        assert saved["Budget"]["C2"].value == 40
        assert saved["Budget"]["G2"].value is None
        assert saved["Budget"]["D2"].value == 9007199254740993
        assert formula["Notes"].sheet_state == "hidden"
        formula.close()
        saved.close()
        cross_reader = {"status": "PASS", "library": "openpyxl", "version": openpyxl.__version__,
                        "scope": "read-only comparison of synthetic before.xlsx; not real Excel save verification"}

    metadata = json.loads(command(["cargo", "metadata", "--offline", "--format-version", "1"], "metadata"))
    packages = [{"name": p["name"], "version": p["version"], "license": p["license"]}
                for p in metadata["packages"] if p["source"]]
    assert all(p["license"] and not any(word in p["license"] for word in ("AGPL", "GPL-", "LGPL")) for p in packages)
    sources = {}
    for p in sorted(HERE.rglob("*")):
        if not p.is_file() or any(s in p.relative_to(HERE).parts for s in ("target", "evidence", "__pycache__", "generated")) or p.name == "Cargo.lock":
            continue
        sources[str(p.relative_to(HERE))] = hashlib.sha256(p.read_bytes()).hexdigest()
    report = {
        "recorded_at_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "scope": "author-run standalone experiment; not independent acceptance or production integration",
        "base_commit": subprocess.check_output(["git", "merge-base", "HEAD", "origin/main"], cwd=REPO, text=True).strip(),
        "environment": {"platform": platform.platform(), "python": platform.python_version(),
                        "rustc": command(["rustc", "--version"], "rustc").strip()},
        "source_sha256": sources, "checks": checks, "fixtures": results,
        "coordinate_comparisons": comparisons, "cross_reader": cross_reader, "dependency_licenses": packages,
        "unrun": ["real Excel GUI/save roundtrip", "OS-enforced worker memory/time isolation", "coverage-guided fuzzing", "independent reviewer acceptance", "Memory/Organize/DTO/Tauri integration"],
    }
    (HERE / "evidence" / "validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"checks": checks, "fixtures_passed": len(results), "comparisons_passed": len(comparisons), "cross_reader": cross_reader}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
